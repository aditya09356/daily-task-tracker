
import sqlite3
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)

APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "tasks.db"

STATUSES = ["Not Started", "In Progress", "Completed", "Blocked"]
PRIORITIES = ["Low", "Medium", "High", "Critical"]

st.set_page_config(
    page_title="Daily Task Tracker",
    page_icon="📋",
    layout="wide",
)


# -------------------- DATABASE --------------------

def get_connection():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_connection() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                work_date TEXT NOT NULL,
                project TEXT NOT NULL,
                task TEXT NOT NULL,
                description TEXT DEFAULT '',
                hours REAL NOT NULL DEFAULT 0,
                priority TEXT NOT NULL DEFAULT 'Medium',
                status TEXT NOT NULL DEFAULT 'Not Started',
                blockers TEXT DEFAULT '',
                next_steps TEXT DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)


def add_task(data):
    with get_connection() as conn:
        conn.execute("""
            INSERT INTO tasks (
                work_date, project, task, description, hours,
                priority, status, blockers, next_steps
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            data["work_date"].isoformat(),
            data["project"].strip(),
            data["task"].strip(),
            data["description"].strip(),
            data["hours"],
            data["priority"],
            data["status"],
            data["blockers"].strip(),
            data["next_steps"].strip(),
        ))


def load_tasks(start_date, end_date):
    with get_connection() as conn:
        rows = conn.execute("""
            SELECT id, work_date, project, task, description, hours,
                   priority, status, blockers, next_steps
            FROM tasks
            WHERE work_date BETWEEN ? AND ?
            ORDER BY work_date DESC, id DESC
        """, (start_date.isoformat(), end_date.isoformat())).fetchall()

    return pd.DataFrame([dict(row) for row in rows])


def delete_task(task_id):
    with get_connection() as conn:
        conn.execute("DELETE FROM tasks WHERE id = ?", (task_id,))

def update_task(task_id, data):
    with get_connection() as conn:
        conn.execute("""
            UPDATE tasks
            SET work_date = ?,
                project = ?,
                task = ?,
                description = ?,
                hours = ?,
                priority = ?,
                status = ?,
                blockers = ?,
                next_steps = ?
            WHERE id = ?
        """, (
            data["work_date"].isoformat(),
            data["project"].strip(),
            data["task"].strip(),
            data["description"].strip(),
            data["hours"],
            data["priority"],
            data["status"],
            data["blockers"].strip(),
            data["next_steps"].strip(),
            task_id,
        ))


init_db()


# -------------------- REPORTING --------------------

def build_summary(df):
    if df.empty:
        return {
            "Total tasks": 0,
            "Completed": 0,
            "In progress": 0,
            "Blocked": 0,
            "Hours logged": 0.0,
        }

    return {
        "Total tasks": len(df),
        "Completed": int((df["status"] == "Completed").sum()),
        "In progress": int((df["status"] == "In Progress").sum()),
        "Blocked": int((df["status"] == "Blocked").sum()),
        "Hours logged": round(float(df["hours"].sum()), 2),
    }


def create_excel(df, start_date, end_date):
    output = BytesIO()
    summary = build_summary(df)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df = pd.DataFrame(
            list(summary.items()), columns=["Metric", "Value"]
        )
        summary_df.to_excel(
            writer, sheet_name="Summary", index=False, startrow=3
        )

        project_summary = (
            df.groupby("project", as_index=False)
            .agg(Tasks=("id", "count"), Hours=("hours", "sum"))
            if not df.empty
            else pd.DataFrame(columns=["project", "Tasks", "Hours"])
        )
        project_summary.to_excel(
            writer, sheet_name="Summary", index=False, startrow=12
        )

        report_df = df.drop(columns=["id"], errors="ignore").copy()
        report_df.to_excel(writer, sheet_name="Task Details", index=False)

        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill(
                    "solid", fgColor="1F4E78"
                )

            for column_cells in sheet.columns:
                max_length = max(
                    len(str(cell.value or ""))
                    for cell in list(column_cells)[:100]
                )
                width = min(max(max_length + 2, 12), 40)
                sheet.column_dimensions[
                    get_column_letter(column_cells[0].column)
                ].width = width

        summary_sheet = writer.book["Summary"]
        summary_sheet["A1"] = "DAILY / WEEKLY WORK REPORT"
        summary_sheet["A2"] = f"Period: {start_date} to {end_date}"
        summary_sheet["A1"].font = Font(size=16, bold=True)

    output.seek(0)
    return output.getvalue()


def create_pdf(df, start_date, end_date):
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F4E78"),
    ))

    story = [
        Paragraph("Daily / Weekly Work Report", styles["ReportTitle"]),
        Paragraph(
            f"Reporting period: {start_date} to {end_date}",
            styles["Normal"],
        ),
        Spacer(1, 6 * mm),
    ]

    summary = build_summary(df)
    metrics = [["Metric", "Value"]] + [
        [str(key), str(value)] for key, value in summary.items()
    ]

    summary_table = Table(metrics, repeatRows=1)
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([summary_table, Spacer(1, 6 * mm)])

    columns = [
        ("work_date", "Date"),
        ("project", "Project"),
        ("task", "Task"),
        ("hours", "Hours"),
        ("priority", "Priority"),
        ("status", "Status"),
        ("blockers", "Blockers"),
        ("next_steps", "Next steps"),
    ]

    table_data = [[label for _, label in columns]]

    for _, row in df.iterrows():
        table_data.append([
            Paragraph(
                str(row.get(key, "") or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
                styles["BodyText"],
            )
            for key, _ in columns
        ])

    if len(table_data) == 1:
        table_data.append(["No tasks recorded"] + [""] * (len(columns) - 1))

    task_table = Table(
        table_data,
        repeatRows=1,
        colWidths=[22*mm, 27*mm, 43*mm, 15*mm,
                   20*mm, 25*mm, 48*mm, 48*mm],
    )
    task_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(task_table)
    doc.build(story)

    output.seek(0)
    return output.getvalue()


# -------------------- USER INTERFACE --------------------

st.title("📋 Daily Task Tracker")
st.caption("Record your work, monitor progress, and prepare manager reports.")

with st.sidebar:
    st.header("Report period")
    today = date.today()
    period = st.radio("Choose period", ["All time", "Daily", "Weekly", "Custom"])

    if period == "All time":
        everything = load_tasks(date(2000, 1, 1), date(2100, 12, 31))
        if everything.empty:
            start_date = end_date = today
        else:
            dates = everything["work_date"].astype(str)
            start_date = date.fromisoformat(dates.min())
            end_date = date.fromisoformat(dates.max())
    elif period == "Daily":
        report_date = st.date_input("Work date", value=today)
        start_date = report_date
        end_date = report_date
    elif period == "Weekly":
        report_date = st.date_input("Any date in the week", value=today)
        start_date = report_date - timedelta(days=report_date.weekday())
        end_date = start_date + timedelta(days=6)
    else:
        start_date = st.date_input("Start date", value=today - timedelta(days=6))
        end_date = st.date_input("End date", value=today)

    st.caption(f"Selected: {start_date} to {end_date}")

tab_entry, tab_report, tab_manage = st.tabs(
    ["➕ Add Task", "📊 Dashboard & Reports", "🗂️ Manage Tasks"]
)

with tab_entry:
    st.subheader("Record a task")

    with st.form("task_form", clear_on_submit=True):
        work_date = st.date_input("Work date", value=date.today())
        project = st.text_input("Project / Workstream")
        task = st.text_input("Task title")
        description = st.text_area("Work performed / details")
        col1, col2, col3 = st.columns(3)

        with col1:
            hours = st.number_input(
                "Hours spent", min_value=0.0, max_value=24.0,
                value=1.0, step=0.5
            )
        with col2:
            priority = st.selectbox("Priority", PRIORITIES, index=1)
        with col3:
            status = st.selectbox("Status", STATUSES, index=1)

        blockers = st.text_area("Blockers / dependencies")
        next_steps = st.text_area("Next steps")

        submitted = st.form_submit_button(
            "Save task", type="primary", use_container_width=True
        )

        if submitted:
            if not project.strip() or not task.strip():
                st.error("Project and task title are required.")
            else:
                add_task({
                    "work_date": work_date,
                    "project": project,
                    "task": task,
                    "description": description,
                    "hours": hours,
                    "priority": priority,
                    "status": status,
                    "blockers": blockers,
                    "next_steps": next_steps,
                })
                st.success("Task saved successfully.")
                st.rerun()

with tab_report:
    st.subheader("Reporting dashboard")
    df = load_tasks(start_date, end_date)

    if not df.empty:
        status_filter = st.multiselect(
            "Filter by status", STATUSES, default=STATUSES
        )
        df = df[df["status"].isin(status_filter)]

    summary = build_summary(df)
    cols = st.columns(5)

    for col, (label, value) in zip(cols, summary.items()):
        col.metric(label, value)

    if df.empty:
        st.info("No tasks found for this reporting period.")
    else:
        st.subheader("Task status")
        status_counts = (
            df["status"].value_counts().reindex(STATUSES, fill_value=0)
        )
        st.bar_chart(status_counts)

        st.subheader("Hours by project")
        project_hours = df.groupby("project")["hours"].sum().sort_values(
            ascending=False
        )
        st.bar_chart(project_hours)

        st.subheader("Task details")
        st.dataframe(
            df.drop(columns=["id"], errors="ignore"),
            use_container_width=True,
            hide_index=True,
        )

    excel_data = create_excel(df, start_date, end_date)
    pdf_data = create_pdf(df, start_date, end_date)

    col1, col2 = st.columns(2)
    filename = f"work_report_{start_date}_{end_date}"

    col1.download_button(
        "⬇️ Download Excel report",
        data=excel_data,
        file_name=f"{filename}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    col2.download_button(
        "⬇️ Download PDF report",
        data=pdf_data,
        file_name=f"{filename}.pdf",
        mime="application/pdf",
        use_container_width=True,
    )

# -------------------- MANAGE TASKS --------------------

STATUS_ICONS = {
    "Not Started": "⚪", "In Progress": "🟡", "Completed": "🟢", "Blocked": "🔴",
}
PRIORITY_ICONS = {"Low": "🔹", "Medium": "🔸", "High": "🔶", "Critical": "🟥"}
PAGE_SIZE = 20


def set_status(task_id, status):
    with get_connection() as conn:
        conn.execute("UPDATE tasks SET status = ? WHERE id = ?", (status, task_id))


@st.dialog("✏️ Edit task", width="large")
def edit_dialog(current):
    task_id = int(current["id"])
    with st.form(f"edit_form_{task_id}", border=False):
        e_date = st.date_input(
            "Work date", value=date.fromisoformat(str(current["work_date"]))
        )
        c1, c2 = st.columns(2)
        e_project = c1.text_input("Project / Workstream", value=current["project"] or "")
        e_task = c2.text_input("Task title", value=current["task"] or "")
        e_description = st.text_area(
            "Work performed / details", value=current["description"] or ""
        )
        c1, c2, c3 = st.columns(3)
        e_hours = c1.number_input(
            "Hours spent", min_value=0.0, max_value=24.0,
            value=float(current["hours"] or 0.0), step=0.5,
        )
        e_priority = c2.selectbox(
            "Priority", PRIORITIES,
            index=PRIORITIES.index(current["priority"])
            if current["priority"] in PRIORITIES else 1,
        )
        e_status = c3.selectbox(
            "Status", STATUSES,
            index=STATUSES.index(current["status"])
            if current["status"] in STATUSES else 1,
        )
        e_blockers = st.text_area(
            "Blockers / dependencies", value=current["blockers"] or ""
        )
        e_next_steps = st.text_area("Next steps", value=current["next_steps"] or "")

        b1, b2 = st.columns([3, 1])
        save = b1.form_submit_button(
            "💾 Save changes", type="primary", use_container_width=True
        )
        cancel = b2.form_submit_button("Cancel", use_container_width=True)

    if cancel:
        st.rerun()
    if save:
        if not e_project.strip() or not e_task.strip():
            st.error("Project and task title are required.")
            return
        update_task(task_id, {
            "work_date": e_date, "project": e_project, "task": e_task,
            "description": e_description, "hours": e_hours,
            "priority": e_priority, "status": e_status,
            "blockers": e_blockers, "next_steps": e_next_steps,
        })
        st.toast("Task updated.", icon="✅")
        st.rerun()


@st.dialog("🗑️ Delete task")
def delete_dialog(current):
    st.warning(
        f"Permanently delete **{current['task']}** "
        f"({current['work_date']})? This cannot be undone."
    )
    c1, c2 = st.columns(2)
    if c1.button("Yes, delete", type="primary", use_container_width=True):
        delete_task(int(current["id"]))
        st.toast("Task deleted.", icon="🗑️")
        st.rerun()
    if c2.button("Cancel", use_container_width=True):
        st.rerun()


def render_task_card(row):
    with st.container(border=True):
        info, actions = st.columns([5, 2], vertical_alignment="center")

        with info:
            st.markdown(f"**{row['task']}**")
            st.caption(
                f"📅 {row['work_date']} · 📁 {row['project']} · "
                f"⏱️ {row['hours']:g} h · "
                f"{STATUS_ICONS.get(row['status'], '')} {row['status']} · "
                f"{PRIORITY_ICONS.get(row['priority'], '')} {row['priority']}"
            )
            if row["description"]:
                st.write(row["description"])
            if row["blockers"]:
                st.markdown(f"🚧 **Blockers:** {row['blockers']}")
            if row["next_steps"]:
                st.markdown(f"➡️ **Next:** {row['next_steps']}")

        with actions:
            b1, b2, b3, b4 = st.columns(4)
            tid = int(row["id"])
            if b1.button("✏️", key=f"edit_{tid}", help="Edit task"):
                edit_dialog(row)
            if row["status"] != "Completed" and b2.button(
                "✅", key=f"done_{tid}", help="Mark as completed"
            ):
                set_status(tid, "Completed")
                st.toast("Marked as completed.", icon="✅")
                st.rerun()
            if b3.button("📄", key=f"dup_{tid}", help="Duplicate to today"):
                add_task({**row.to_dict(), "work_date": date.today(),
                          "status": "Not Started"})
                st.toast("Task duplicated for today.", icon="📄")
                st.rerun()
            if b4.button("🗑️", key=f"del_{tid}", help="Delete task"):
                delete_dialog(row)


with tab_manage:
    st.subheader("Manage saved tasks")
    all_tasks = load_tasks(date(2000, 1, 1), date(2100, 12, 31))

    if all_tasks.empty:
        st.info("No tasks have been saved yet.")
    else:
        f1, f2, f3 = st.columns([3, 2, 2])
        query = f1.text_input(
            "🔍 Search", placeholder="Search task, project, details…"
        )
        status_sel = f2.multiselect("Status", STATUSES, placeholder="All statuses")
        project_sel = f3.multiselect(
            "Project", sorted(all_tasks["project"].dropna().unique()),
            placeholder="All projects",
        )

        shown = all_tasks
        if query.strip():
            text_cols = ["task", "project", "description", "blockers", "next_steps"]
            haystack = shown[text_cols].fillna("").agg(" ".join, axis=1)
            shown = shown[haystack.str.contains(query.strip(), case=False, regex=False)]
        if status_sel:
            shown = shown[shown["status"].isin(status_sel)]
        if project_sel:
            shown = shown[shown["project"].isin(project_sel)]

        st.caption(f"Showing {len(shown)} of {len(all_tasks)} tasks")

        if shown.empty:
            st.info("No tasks match your filters.")
        else:
            limit = st.session_state.get("manage_limit", PAGE_SIZE)
            for _, row in shown.head(limit).iterrows():
                render_task_card(row)
            if len(shown) > limit:
                if st.button(f"Show more ({len(shown) - limit} remaining)"):
                    st.session_state["manage_limit"] = limit + PAGE_SIZE
                    st.rerun()

st.divider()
st.caption("Local database: tasks.db · Back up this file regularly.")
